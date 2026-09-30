"""
Build the Aeropanel opportunity workbook - Case Study B
=======================================================
Combines the latest DIBBS collector run, DIBBS package-status checks, parsed PDFs and the latest
SAM.gov run into one workbook: management summary first, then the evidence-linked G-only line list
and every separate population.

    python build_report.py
Output: data/reports/Aeropanel_Opportunity_Feed_<timestamp>.xlsx  (+ CSV copy of the main list)

Status rules (documented in README, decisions 17 and 19-22):
  Excluded (awarded/cancelled) : SAM.gov award notice for the solicitation, or DIBBS package status Awarded/Cancelled
  Currently open              : DIBBS package status Open and Return By >= today  | SAM.gov active notice, deadline >= today
  Late route verified         : DIBBS package status Open (quoting enabled) although Return By has passed
  Follow-up needed            : past deadline, no award found in the checks performed, late quote unconfirmed
  Status unresolved           : package page not checked / Removed / unparsed / contradictory
"""
import json, re, sys
from datetime import datetime, timezone
from pathlib import Path
import pandas as pd
from common import DATA, CUTOFF, CATEGORY_NAMES, norm_sol, norm_nsn, utc_stamp, log_run

# Treat DIBBS "Open" after the return-by date as a verified late-quotation route.
# DIBBS defines Open as "RFQs available for quoting". Set to False to downgrade these to Follow-up needed.
LATE_ROUTE_FROM_OPEN_STATUS = True

TODAY = pd.Timestamp.now().normalize()
NO_AWARD_FOUND = ["Currently open", "Late route verified", "Follow-up needed"]
STATUS_ORDER = NO_AWARD_FOUND + ["Status unresolved"]
GROUPS = ["core", "core-adjacent", "adjacent", "INSCO"]


# ------------------------------------------------------------------ loading helpers
def read(p: Path) -> pd.DataFrame:
    if p.exists() and p.stat().st_size > 2:
        return pd.read_csv(p, dtype=str, low_memory=False)
    return pd.DataFrame()


def latest(pattern: str, must_have: str) -> Path | None:
    runs = sorted(p for p in (DATA / "runs").glob(pattern) if (p / must_have).exists())
    return runs[-1] if runs else None


def num(x):
    try:
        return float(str(x).replace(",", ""))
    except (TypeError, ValueError):
        return None


def d(x):
    return pd.to_datetime(x, errors="coerce", format="%m-%d-%Y")


# ------------------------------------------------------------------ DIBBS package status
def load_package_status() -> pd.DataFrame:
    frames = [read(p) for p in sorted((DATA / "dibbs_status").glob("*/package_status.csv"))]
    frames = [f for f in frames if len(f)]
    if not frames:
        return pd.DataFrame(columns=["sol_key"])
    df = pd.concat(frames, ignore_index=True).sort_values("checked_utc")
    return df.drop_duplicates("sol_key", keep="last").set_index("sol_key")


def load_pdf_extracts():
    folders = sorted((DATA / "pdf_extracts").glob("pdfparse_*"))
    summ, lines = pd.DataFrame(), pd.DataFrame()
    for f in folders:                                   # later parses win
        s, l = read(f / "pdf_summary.csv"), read(f / "pdf_lines.csv")
        if len(s):
            summ = pd.concat([summ, s]).drop_duplicates("sol_key", keep="last")
        if len(l):
            lines = pd.concat([lines, l]).drop_duplicates("line_key", keep="last")
    return summ.set_index("sol_key") if len(summ) else summ, lines


def amsc_lookup() -> dict:
    """NSN -> (AMSC, source) from every DIBBS listing collected (AMSC is a property of the item/NSN)."""
    out = {}
    for f in sorted((DATA / "runs").glob("dibbs_*/dibbs_all_rows.csv")):
        df = read(f)
        if not len(df) or "nsn" not in df:
            continue
        for r in df[["nsn", "amsc", "retrieved_utc"]].dropna(subset=["nsn"]).itertuples(index=False):
            if str(r.amsc or "").strip():
                out[r.nsn] = (str(r.amsc).strip().upper(), f"DIBBS Text Search listing for NSN {r.nsn}, retrieved {r.retrieved_utc}")
    return out


# ------------------------------------------------------------------ status rules
def dibbs_status(sol_key, return_by_text, pkg: pd.DataFrame, sam_awards: dict, retrieved):
    if sol_key in sam_awards:
        a = sam_awards[sol_key]
        return "Excluded: awarded", f"SAM.gov award notice {a['notice']} (award {a['award']}, {a['date']})", ""
    if sol_key not in pkg.index:
        return ("Status unresolved",
                f"Solicitation-level status check not yet performed. Listed in DIBBS open-document Text Search on {retrieved}.", "")
    p = pkg.loc[sol_key]
    st, checked = str(p.get("status") or ""), p.get("checked_utc", "")
    rb = d(p.get("return_by")) if str(p.get("return_by") or "") else d(return_by_text)
    ev = f"DIBBS package page Status '{p.get('status_raw', st)}', Return By {p.get('return_by', return_by_text)}, checked {checked} (saved: {p.get('evidence_file', '')})"
    if st == "Open":
        if pd.notna(rb) and rb >= TODAY:
            return "Currently open", ev, checked
        if LATE_ROUTE_FROM_OPEN_STATUS:
            return ("Late route verified",
                    ev + ". Return-by date passed but DIBBS still shows Status Open with quoting enabled "
                         "(DIBBS: Open = 'RFQs available for quoting'). Condition: DLA may still award to another timely offer.",
                    checked)
        return "Follow-up needed", ev + ". Late quote acceptance unconfirmed.", checked
    if st in ("Awarded", "Cancelled"):
        return f"Excluded: {st.lower()}", ev, checked
    if st == "Closed":
        return "Follow-up needed", ev + ". No award found; closed to quoting.", checked
    if st == "Removed":
        return "Status unresolved", ev + ". Removed from DIBBS; replacement solicitation (by PR number) not checked.", checked
    return "Status unresolved", ev + f". Could not interpret status ({p.get('parse_issue', '')}).", checked


def sam_status(r, sam_awards):
    if r["sol_key"] and r["sol_key"] in sam_awards:
        a = sam_awards[r["sol_key"]]
        return "Excluded: awarded", f"SAM.gov award notice {a['notice']} (award {a['award']}, {a['date']})", ""
    dl = pd.to_datetime(r.get("ResponseDeadLine"), errors="coerce", utc=True)
    snap = r.get("snapshot_id", "")
    active = str(r.get("Active", "")).strip().upper() == "YES"
    if pd.notna(dl) and dl.tz_localize(None) >= TODAY and active:
        return "Currently open", f"SAM.gov notice active, response deadline {r.get('ResponseDeadLine')} (snapshot {snap})", snap
    if pd.notna(dl):
        return ("Follow-up needed", f"Deadline {r.get('ResponseDeadLine')} passed; no award notice for this solicitation in "
                                    f"SAM.gov snapshot {snap}. Late quote acceptance unconfirmed.", snap)
    return "Status unresolved", f"No usable response deadline in SAM.gov snapshot {snap}", snap


# ------------------------------------------------------------------ build line tables
def build_dibbs_lines(g: pd.DataFrame, population: str, pkg, pdf_summ, pdf_lines, sam_awards, sam_by_sol):
    out = []
    for _, r in g.fillna("").iterrows():
        sk, nsn = r["sol_key"], r.get("nsn", "")
        status, evidence, checked = dibbs_status(sk, r.get("return_by"), pkg, sam_awards, r.get("retrieved_utc"))
        s = pdf_summ.loc[sk] if len(pdf_summ) and sk in pdf_summ.index else None
        base = {
            "population": population, "source": "DIBBS", "solicitation": r.get("solicitation"), "sol_key": sk,
            "official_link": r.get("official_link"), "pdf_link": r.get("pdf_url"), "nsn": nsn,
            "item_description": r.get("nomenclature"), "category": r.get("category"),
            "category_name": r.get("category_name"), "group": r.get("group"), "primary_keyword": r.get("primary_keyword"),
            "keyword_tags": r.get("keyword_tags"), "review_flag": r.get("review_flag"),
            "amsc": r.get("amsc"), "amsc_source": r.get("amsc_source"),
            "published": r.get("issued"), "published_basis": "DIBBS Issued date", "deadline": r.get("return_by"),
            "status": status, "status_evidence": evidence, "status_checked": checked,
            "also_on_sam": "; ".join(sam_by_sol.get(sk, [])), "retrieved_utc": r.get("retrieved_utc"),
            "found_by_search_terms": r.get("found_by_terms"),
        }
        if s is not None:
            base.update({"set_aside": s.get("set_aside"),
                         "gates": "; ".join(k.replace("_", " ") for k in ["first_article_required", "cmmc_required",
                                            "export_controlled_data", "critical_application_item"] if str(s.get(k)) == "True"),
                         "pdf_file": s.get("pdf_file")})
        # quantities: PDF CLINs (line level) when available, else the Text Search quantity
        pl = pdf_lines[(pdf_lines["sol_key"] == sk) & (pdf_lines["nsn"] == nsn) &
                       (pdf_lines["is_first_article_line"].astype(str) != "True")] if len(pdf_lines) else pd.DataFrame()
        lines = ([{"line_id": l["line_key"], "quantity": num(l["quantity"]), "unit": l["unit"], "pr": l["pr"],
                   "qty_source": f"{l['source']} (CLIN {l['clin']})"} for _, l in pl.iterrows()]
                 if len(pl) else
                 [{"line_id": f"DIBBS|{sk}|{nsn}", "quantity": num(r.get("quantity_listed")), "unit": "",
                   "pr": r.get("purchase_requests"),
                   "qty_source": "DIBBS Text Search result (may combine several PR lines; PDF not parsed)"}])
        for ln in lines:
            row = {**base, **ln}
            price, hist_nsn = None, ""
            if s is not None:
                price = num(s.get("hist_unit_price"))
            if price:
                row.update({"unit_price": price, "price_basis": "Exact-NSN procurement history, most recent award",
                            "price_date": s.get("hist_award_date"), "price_age_years": num(s.get("hist_age_years")),
                            "price_source": f"{s.get('hist_source')}; contract {s.get('hist_contract')}, CAGE {s.get('hist_cage')}, qty {s.get('hist_order_qty')}",
                            "price_note": s.get("price_note")})
            else:
                row.update({"unit_price": None, "price_basis": "Unpriced",
                            "price_note": ("PDF not parsed yet" if s is None else s.get("price_note"))})
            row["estimated_value"] = (row["quantity"] * price) if (price and row["quantity"]) else None
            out.append(row)
    return out


def build_sam_lines(sam: pd.DataFrame, population: str, amsc: dict, dibbs_sols: set, sam_awards, pdf_summ):
    g, review, nong, crossposted = [], [], [], []
    for _, r in sam.fillna("").iterrows():
        sk = r.get("sol_key") or norm_sol(r.get("Sol#"))
        r = r.copy(); r["sol_key"] = sk
        if sk and sk in dibbs_sols:
            crossposted.append({"NoticeId": r.get("NoticeId"), "solicitation": r.get("Sol#"), "sol_key": sk,
                                "title": r.get("Title"), "resolution": "Cross-posted: counted once, under the DIBBS record"})
            continue
        nsn = r.get("nsn_found") or ""
        code, src = amsc.get(nsn, ("", ""))
        status, evidence, checked = sam_status(r, sam_awards)
        row = {"population": population, "source": "SAM.gov", "solicitation": r.get("Sol#"), "sol_key": sk,
               "official_link": r.get("public_link"), "notice_id": r.get("NoticeId"), "nsn": nsn,
               "item_description": r.get("Title"), "category": r.get("category"), "category_name": r.get("category_name"),
               "group": r.get("group"), "primary_keyword": r.get("primary_keyword"), "keyword_tags": r.get("keyword_tags"),
               "review_flag": r.get("review_flag"), "amsc": code, "amsc_source": src or "No AMSC in SAM.gov; NSN not found on DIBBS",
               "published": r.get("PostedDate"), "published_basis": "SAM.gov PostedDate", "deadline": r.get("ResponseDeadLine"),
               "status": status, "status_evidence": evidence, "status_checked": checked, "set_aside": r.get("SetASide"),
               "retrieved_utc": r.get("snapshot_id"), "line_id": f"SAM|{r.get('NoticeId')}",
               "quantity": None, "qty_source": "Not in SAM.gov structured data (see notice/attachments)",
               "unit_price": None, "price_basis": "Unpriced", "estimated_value": None,
               "price_note": "Quantity not available in SAM.gov data"}
        if code == "G":
            g.append(row)
        elif code:
            row["reason"] = f"AMSC {code} (not G)"; nong.append(row)
        else:
            row["reason"] = "AMSC unverified: SAM.gov has no AMSC field" + ("" if nsn else " and no NSN in the notice")
            review.append(row)
    return g, review, nong, crossposted


# ------------------------------------------------------------------ helper flags for formulas
def add_flags(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["no_award_found"] = df["status"].isin(NO_AWARD_FOUND).astype(int)
    df["unpriced"] = df["estimated_value"].isna().astype(int)
    df["f_sol_cat_status"] = (~df.duplicated(["sol_key", "category", "status"])).astype(int)
    df["f_sol_status"] = (~df.duplicated(["sol_key", "status"])).astype(int)
    nf = df["no_award_found"] == 1
    df["f_sol_cat_nf"] = 0; df["f_sol_nf"] = 0; df["f_sol_group_nf"] = 0
    df.loc[nf, "f_sol_cat_nf"] = (~df[nf].duplicated(["sol_key", "category"])).astype(int)
    df.loc[nf, "f_sol_nf"] = (~df[nf].duplicated(["sol_key"])).astype(int)
    df.loc[nf, "f_sol_group_nf"] = (~df[nf].duplicated(["sol_key", "group"])).astype(int)
    return df


# ------------------------------------------------------------------ workbook
def write_workbook(path: Path, main: pd.DataFrame, sheets: dict, meta: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    F = "Arial"
    bold, head_fill = Font(name=F, bold=True), PatternFill("solid", fgColor="DDE4EE")
    wb = Workbook()
    ws = wb.active; ws.title = "Summary"

    # ---- G_Main first (formulas on Summary reference it)
    cols = list(main.columns)
    gm = wb.create_sheet("G_Main")
    gm.append(cols)
    for row in main.itertuples(index=False):
        gm.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    n = len(main) + 1
    rng = lambda c: f"G_Main!${L[c]}$2:${L[c]}${max(n, 2)}"

    def fmt_sheet(sh, df_cols, money=(), link_cols=()):
        for c in sh[1]:
            c.font, c.fill = bold, head_fill
        sh.freeze_panes = "A2"
        if sh.max_row > 1:
            sh.auto_filter.ref = sh.dimensions
        for i, c in enumerate(df_cols, 1):
            sh.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(str(c)) + 2))
            if c in money:
                for cell in sh[get_column_letter(i)][1:]:
                    cell.number_format = '"$"#,##0.00'
            if c in link_cols:
                for cell in sh[get_column_letter(i)][1:]:
                    if isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value; cell.font = Font(name=F, color="0563C1", underline="single")
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if cell.font.color is None or cell.font.color.rgb != "FF0563C1":
                    cell.font = Font(name=F)

    fmt_sheet(gm, cols, money=("unit_price", "estimated_value"), link_cols=("official_link", "pdf_link"))

    # ---- Summary
    def put(r, c, v, b=False, money=False):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, bold=b)
        if money:
            cell.number_format = '"$"#,##0'
        return cell

    r = 1
    put(r, 1, "Aeropanel Opportunity Feed: DIBBS + SAM.gov, 2026 AMSC G opportunities", True).font = Font(name=F, bold=True, size=14)
    for k, v in meta["header"]:
        r += 1; put(r, 1, k, True); put(r, 2, v)
    r += 2; put(r, 1, "Headline results (confirmed AMSC G, 2026 main population only)", True)
    heads = [("(1) Unique solicitations, no award found in stated checks", f"=COUNTIFS({rng('f_sol_nf')},1)",
              f"=SUMIFS({rng('estimated_value')},{rng('no_award_found')},1)"),
             ("(2) Currently open (subset)", f'=COUNTIFS({rng("status")},"Currently open",{rng("f_sol_status")},1)',
              f'=SUMIFS({rng("estimated_value")},{rng("status")},"Currently open")'),
             ("(3) Past deadline, late route verified", f'=COUNTIFS({rng("status")},"Late route verified",{rng("f_sol_status")},1)',
              f'=SUMIFS({rng("estimated_value")},{rng("status")},"Late route verified")'),
             ("(4) Past deadline, follow-up leads", f'=COUNTIFS({rng("status")},"Follow-up needed",{rng("f_sol_status")},1)',
              f'=SUMIFS({rng("estimated_value")},{rng("status")},"Follow-up needed")'),
             ("(5) Unpriced lines (no award found)", f"=COUNTIFS({rng('unpriced')},1,{rng('no_award_found')},1)", None),
             ("Outside totals: status unresolved (unique solicitations)", f'=COUNTIFS({rng("status")},"Status unresolved",{rng("f_sol_status")},1)',
              f'=SUMIFS({rng("estimated_value")},{rng("status")},"Status unresolved")')]
    r += 1; put(r, 1, "Measure", True); put(r, 2, "Count", True); put(r, 3, "Estimated value", True)
    for label, cnt, val in heads:
        r += 1; put(r, 1, label); put(r, 2, cnt)
        if val:
            put(r, 3, val, money=True)

    r += 2; put(r, 1, "Category and status summary (brief layout)", True)
    r += 1
    for i, h in enumerate(["Primary category", "Status", "Unique solicitations", "Line items", "Estimated value", "Unpriced lines"], 1):
        put(r, i, h, True).fill = head_fill
    labels = {"C01": "C01 Control / HMI panels"}
    rows = [(f"{c} {CATEGORY_NAMES[c]}", c) for c in CATEGORY_NAMES] + [("Overall, independently deduplicated", None)]
    for label, cat in rows:
        for st in ["All no-award-found"] + NO_AWARD_FOUND:
            r += 1; put(r, 1, label); put(r, 2, st)
            cc = f',{rng("category")},"{cat}"' if cat else ""
            if st == "All no-award-found":
                flag = "f_sol_cat_nf" if cat else "f_sol_nf"
                put(r, 3, f"=COUNTIFS({rng(flag)},1{cc})")
                put(r, 4, f"=COUNTIFS({rng('no_award_found')},1{cc})")
                put(r, 5, f"=SUMIFS({rng('estimated_value')},{rng('no_award_found')},1{cc})", money=True)
                put(r, 6, f"=COUNTIFS({rng('no_award_found')},1,{rng('unpriced')},1{cc})")
            else:
                flag = "f_sol_cat_status" if cat else "f_sol_status"
                sc = f',{rng("status")},"{st}"'
                put(r, 3, f"=COUNTIFS({rng(flag)},1{sc}{cc})")
                put(r, 4, f'=COUNTIFS({rng("status")},"{st}"{cc})')
                put(r, 5, f'=SUMIFS({rng("estimated_value")}{sc}{cc})', money=True)
                put(r, 6, f"=COUNTIFS({rng('unpriced')},1{sc}{cc})")
    r += 1; put(r, 1, "Unique-solicitation counts are recomputed from distinct IDs (helper flags on G_Main); they are not added across rows. "
                      "'All no-award-found' is not added to its status subsets.")

    r += 2; put(r, 1, "Group views (no award found)", True)
    r += 1
    for i, h in enumerate(["Group", "Unique solicitations", "Line items", "Estimated value", "Unpriced lines"], 1):
        put(r, i, h, True).fill = head_fill
    for gname in GROUPS:
        r += 1; put(r, 1, gname)
        gc = f',{rng("group")},"{gname}"'
        put(r, 2, f"=COUNTIFS({rng('f_sol_group_nf')},1{gc})")
        put(r, 3, f"=COUNTIFS({rng('no_award_found')},1{gc})")
        put(r, 4, f"=SUMIFS({rng('estimated_value')},{rng('no_award_found')},1{gc})", money=True)
        put(r, 5, f"=COUNTIFS({rng('no_award_found')},1,{rng('unpriced')},1{gc})")

    r += 2; put(r, 1, "Separate populations (outside the totals above)", True)
    for k, v in meta["populations"]:
        r += 1; put(r, 1, k); put(r, 2, v)
    r += 2; put(r, 1, "Coverage and limitations", True)
    for line in meta["coverage"]:
        r += 1; put(r, 1, line)
    ws.column_dimensions["A"].width = 62; ws.column_dimensions["B"].width = 22
    for col in "CDEF":
        ws.column_dimensions[col].width = 18

    # ---- other sheets
    for name, (df, money, links) in sheets.items():
        sh = wb.create_sheet(name[:31])
        if df is None or not len(df):
            sh.append(["(no records)"]); continue
        sh.append(list(df.columns))
        for row in df.itertuples(index=False):
            sh.append([None if (isinstance(v, float) and pd.isna(v)) else (str(v) if isinstance(v, (list, dict)) else v) for v in row])
        fmt_sheet(sh, list(df.columns), money=money, link_cols=links)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ manual check sample
def manual_sample(main: pd.DataFrame) -> pd.DataFrame:
    picks = []
    nf = main[main["no_award_found"] == 1]
    picks.append(nf.sort_values("estimated_value", ascending=False).head(10))
    for st, k in [("Currently open", 4), ("Late route verified", 4), ("Follow-up needed", 3), ("Status unresolved", 3)]:
        picks.append(main[main["status"] == st].sample(min(k, (main["status"] == st).sum()), random_state=1))
    picks.append(main[main["unpriced"] == 1].head(2))
    picks.append(main[main["source"] == "SAM.gov"].head(2))
    picks.append(main[main["sol_key"] == "SPE7M926T0038"])
    s = pd.concat(picks).drop_duplicates("line_id").head(30)
    s = s[["line_id", "source", "solicitation", "official_link", "nsn", "item_description", "category", "status",
           "quantity", "unit_price", "estimated_value"]].copy()
    for c in ["Keyword/category correct? (Y/N)", "AMSC G confirmed on source? (Y/N)", "Quantity matches source? (Y/N)",
              "Status matches source today? (Y/N)", "Price matches PDF history? (Y/N)", "Checked by", "Check date (UTC)",
              "Result (Pass / Issue)", "Notes"]:
        s[c] = ""
    example = {c: "" for c in s.columns}
    example.update({"line_id": "EXAMPLE ROW - delete", "Keyword/category correct? (Y/N)": "Y",
                    "AMSC G confirmed on source? (Y/N)": "Y", "Quantity matches source? (Y/N)": "Y",
                    "Status matches source today? (Y/N)": "Y", "Price matches PDF history? (Y/N)": "N",
                    "Checked by": "Veera", "Check date (UTC)": "2026-10-01",
                    "Result (Pass / Issue)": "Issue", "Notes": "PDF history price is from 2019; flagged as stale"})
    legend = {c: "" for c in s.columns}
    legend["line_id"] = "LEGEND: fill only the columns from 'Keyword/category correct?' to 'Notes'. Open official_link to check."
    return pd.concat([pd.DataFrame([legend, example]), s], ignore_index=True)


# ------------------------------------------------------------------ main
def main():
    run_id = f"report_{utc_stamp()}"
    counts = {"run_id": run_id, "source": "report"}
    try:
        drun = latest("dibbs_*", "g_candidates_main_2026.csv")
        srun = latest("sam_*", "keyword_matches.csv")
        if not drun and not srun:
            raise RuntimeError("No DIBBS or SAM.gov run outputs found under data/runs")
        dstage = json.loads((drun / "stage_counts.json").read_text()) if drun and (drun / "stage_counts.json").exists() else {}
        sstage = json.loads((srun / "stage_counts.json").read_text()) if srun and (srun / "stage_counts.json").exists() else {}
        pkg = load_package_status()
        pdf_summ, pdf_lines = load_pdf_extracts()
        amsc = amsc_lookup()

        dread = (lambda f: read(drun / f)) if drun else (lambda f: pd.DataFrame())
        sread = (lambda f: read(srun / f)) if srun else (lambda f: pd.DataFrame())

        # SAM.gov award notices by solicitation number
        sam_awards = {}
        aw = sread("award_notices.csv")
        for _, a in aw.iterrows():
            k = norm_sol(a.get("Sol#"))
            if k:
                sam_awards[k] = {"notice": a.get("NoticeId"), "award": a.get("AwardNumber"), "date": a.get("AwardDate")}
        sam_main = sread("keyword_matches.csv")
        sam_by_sol = {}
        for _, s in sam_main.iterrows():
            sam_by_sol.setdefault(s.get("sol_key"), []).append(str(s.get("NoticeId")))

        # DIBBS G lines
        dl = []
        for f, pop in [("g_candidates_main_2026.csv", "main_2026"), ("g_pre2026_carry_forward.csv", "pre2026_carry_forward"),
                       ("g_date_unresolved.csv", "date_unresolved")]:
            dl += build_dibbs_lines(dread(f), pop, pkg, pdf_summ, pdf_lines, sam_awards, sam_by_sol)
        dibbs_lines = pd.DataFrame(dl)
        dibbs_sols = set(dibbs_lines["sol_key"]) if len(dibbs_lines) else set()
        for extra in ["non_g_exclusions.csv", "review_queue_amsc.csv"]:
            e = dread(extra)
            if len(e):
                dibbs_sols |= set(e["sol_key"])

        # SAM lines
        sg, sreview, snong, xpost = build_sam_lines(sam_main, "main_2026", amsc, dibbs_sols, sam_awards, pdf_summ)
        sp_g, sp_rev, sp_nong, sp_x = build_sam_lines(sread("pre2026_still_open.csv"), "pre2026_carry_forward",
                                                      amsc, dibbs_sols, sam_awards, pdf_summ)
        sam_g = pd.DataFrame(sg + sp_g)

        all_g = pd.concat([dibbs_lines, sam_g], ignore_index=True)
        excluded_status = all_g[all_g["status"].str.startswith("Excluded")] if len(all_g) else all_g
        main_pop = all_g[(all_g["population"] == "main_2026") & ~all_g["status"].str.startswith("Excluded")]
        main_pop = main_pop.sort_values(["estimated_value", "quantity"], ascending=[False, False], na_position="last")
        main_pop = add_flags(main_pop.reset_index(drop=True))
        carry = all_g[all_g["population"] == "pre2026_carry_forward"]
        date_unres = all_g[all_g["population"] == "date_unresolved"]

        # changes vs previous DIBBS run (line level)
        changes = pd.DataFrame()
        prev_runs = sorted(p for p in (DATA / "runs").glob("dibbs_*") if (p / "g_candidates_main_2026.csv").exists())
        if drun and len(prev_runs) >= 2:
            prev = read(prev_runs[-2] / "g_candidates_main_2026.csv"); cur = dread("g_candidates_main_2026.csv")
            pk, ck = set(prev.get("line_key", [])), set(cur.get("line_key", []))
            changes = pd.DataFrame([{"line_key": k, "change": "new since previous DIBBS run"} for k in ck - pk] +
                                   [{"line_key": k, "change": "no longer listed (check status)"} for k in pk - ck])
            if len(changes):
                changes["compared"] = f"{prev_runs[-2].name} -> {drun.name}"
        sam_changes = sread("changes_vs_previous.csv")

        # separate populations
        review = pd.concat([dread("review_queue_amsc.csv"), pd.DataFrame(sreview + sp_rev)], ignore_index=True)
        nong = pd.concat([dread("non_g_exclusions.csv"), pd.DataFrame(snong + sp_nong)], ignore_index=True)
        excl = pd.concat([dread("exclusion_log.csv").assign(log="keyword rule"),
                          dread("broad_context_review.csv").assign(log="broad keyword not the item noun"),
                          dread("first_article_lines.csv").assign(log="first-article / test line"),
                          sread("exclusion_log.csv").assign(log="SAM.gov keyword rule"),
                          excluded_status.assign(log="awarded / cancelled"),
                          pd.DataFrame(xpost + sp_x).assign(log="cross-post dedupe")], ignore_index=True)
        text_only = dread("text_hit_not_in_nomenclature.csv")
        desc_only = sread("description_only_review.csv")
        terms = dread("term_summary.csv")

        # meta for Summary
        pkg_checked = main_pop["status_checked"].astype(str).str.len().gt(0).sum() if len(main_pop) else 0
        header = [("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
                  ("Reporting window", f"2026-01-01 to {CUTOFF} ET (PROVISIONAL cutoff: common cutoff not yet confirmed by Jonah)"),
                  ("DIBBS run", f"{drun.name if drun else 'none'}: {dstage.get('terms', '?')} search terms, status {dstage.get('status', '?')}"),
                  ("SAM.gov run", f"{srun.name if srun else 'none'}: {sstage.get('retrieval', '?')}, snapshot "
                                  f"{', '.join(s['snapshot_id'] for s in sstage.get('snapshots', []))}"),
                  ("Status checks", f"DIBBS package pages checked for {len(pkg)} solicitations; SAM.gov award notices "
                                    f"({len(sam_awards)} solicitation numbers) matched; DIBBS Awards database not queried per solicitation"),
                  ("Main-list lines with a solicitation-level status check (DIBBS package page or SAM.gov snapshot)", f"{pkg_checked} of {len(main_pop)}"),
                  ("Late-route rule", "DIBBS package Status 'Open' after Return By = late route verified" if LATE_ROUTE_FROM_OPEN_STATUS
                   else "Past-deadline items are follow-up leads only")]
        pops = [("Pre-2026 carry-forward (still open)", len(carry)), ("Date unresolved", len(date_unres)),
                ("AMSC missing / unverified (review queue)", len(review)), ("Non-G exclusions", len(nong)),
                ("Excluded by rules / awarded / cross-post / test lines", len(excl)),
                ("Broad keyword not the item noun", len(dread("broad_context_review.csv"))),
                ("DIBBS text hits not in the item name", len(text_only)), ("SAM.gov description-only matches", len(desc_only))]
        incomplete = dstage.get("terms_incomplete", []) if dstage else []
        coverage = [f"DIBBS: RFQ Solicitation Text Search (open-document index). Terms capped at {dstage.get('max_pages', '?')} pages x 100: "
                    f"{', '.join(incomplete) or 'none'} (DIBBS reports at most 5,000 hits per search).",
                    "DIBBS: solicitations removed from DIBBS drop out of the search index; past-deadline coverage is therefore partial.",
                    "SAM.gov: bulk Contract Opportunities CSV (active notices). Archived notices not yet added. SAM.gov API returned empty 404 (documented blocker).",
                    "AMSC: from DIBBS listings only. SAM.gov-only notices without a DIBBS-listed NSN stay in the review queue.",
                    "Prices: most recent exact-NSN award in the solicitation PDF. DLA management price not sourced; lines without history are unpriced (never $0).",
                    "Quantities: PDF line items (CLINs) where parsed; otherwise the DIBBS listing quantity (may combine PR lines)."]
        meta = {"header": header, "populations": pops, "coverage": coverage}

        money, links = ("unit_price", "estimated_value"), ("official_link", "pdf_link", "public_link")
        runlog = pd.DataFrame([json.loads(l) for l in (DATA / "run_log.jsonl").read_text().splitlines() if l.strip()]) \
            if (DATA / "run_log.jsonl").exists() else pd.DataFrame()
        sheets = {"Review_AMSC": (review, money, links), "NonG_Exclusions": (nong, money, links),
                  "Exclusion_Log": (excl, money, links), "Carry_Forward_pre2026": (carry, money, links),
                  "Date_Unresolved": (date_unres, money, links), "SAM_Description_Only": (desc_only, (), ("public_link",)),
                  "DIBBS_Text_Hit_Only": (text_only, (), ()), "Changes_DIBBS": (changes, (), ()),
                  "Changes_SAM": (sam_changes, (), ()), "Coverage_Terms": (terms, (), ()),
                  "Manual_Checks": (manual_sample(main_pop) if len(main_pop) else pd.DataFrame(), money, links),
                  "Run_Log": (runlog.astype(str) if len(runlog) else runlog, (), ())}
        out = DATA / "reports" / f"Aeropanel_Opportunity_Feed_{utc_stamp()}.xlsx"
        write_workbook(out, main_pop, sheets, meta)
        main_pop.to_csv(out.with_suffix(".main_list.csv"), index=False)

        # NSNs from SAM.gov matches that DIBBS could not classify -> optional lookup list
        need = sorted({r["nsn"] for r in sreview + sp_rev if r.get("nsn")})
        (DATA / "reports" / "sam_nsns_to_check.txt").write_text("\n".join(need))

        nf = main_pop[main_pop["no_award_found"] == 1] if len(main_pop) else main_pop
        counts.update({"workbook": str(out), "main_lines": len(main_pop),
                       "status_counts": main_pop["status"].value_counts().to_dict() if len(main_pop) else {},
                       "no_award_found_unique_solicitations": int(nf["sol_key"].nunique()) if len(nf) else 0,
                       "no_award_found_value": float(nf["estimated_value"].sum()) if len(nf) else 0,
                       "unpriced_lines": int(nf["unpriced"].sum()) if len(nf) else 0,
                       "package_status_checked": len(pkg), "pdf_summaries": len(pdf_summ),
                       "sam_nsns_to_check": len(need), "status": "SUCCESS"})
        log_run(counts)
        print(json.dumps(counts, indent=2, default=str))
    except Exception as e:
        counts.update({"status": "FAILED", "error": f"{type(e).__name__}: {e}"})
        log_run(counts)
        print("\n*** REPORT FAILED ***\n" + counts["error"])
        raise


if __name__ == "__main__":
    main()